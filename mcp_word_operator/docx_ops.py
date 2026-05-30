from __future__ import annotations

import os
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

TABLE_ALIGNMENT_MAP = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}


def _resolve_unit(value: float, unit: str = "pt"):
    if unit == "cm":
        return Cm(value)
    elif unit == "inch":
        return Inches(value)
    else:
        return Pt(value)


def read_document_info(file_path: str) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    paragraphs_info = []
    for i, p in enumerate(doc.paragraphs):
        runs_info = []
        for r in p.runs:
            font = r.font
            runs_info.append({
                "text": r.text,
                "bold": font.bold,
                "italic": font.italic,
                "underline": font.underline,
                "font_name": font.name,
                "font_size": str(font.size) if font.size else None,
                "font_color": str(font.color.rgb) if font.color and font.color.rgb else None,
            })
        paragraphs_info.append({
            "index": i,
            "text": p.text[:200],
            "style": p.style.name if p.style else None,
            "alignment": str(p.alignment) if p.alignment else None,
            "runs_count": len(p.runs),
            "runs": runs_info[:5],
        })

    tables_info = []
    for i, table in enumerate(doc.tables):
        rows_info = []
        for row in table.rows:
            cells = [cell.text[:50] for cell in row.cells]
            rows_info.append(cells)
        tables_info.append({
            "index": i,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "first_rows": rows_info[:3],
        })

    sections_info = []
    for i, section in enumerate(doc.sections):
        sections_info.append({
            "index": i,
            "page_width": str(section.page_width) if section.page_width else None,
            "page_height": str(section.page_height) if section.page_height else None,
            "top_margin": str(section.top_margin) if section.top_margin else None,
            "bottom_margin": str(section.bottom_margin) if section.bottom_margin else None,
            "left_margin": str(section.left_margin) if section.left_margin else None,
            "right_margin": str(section.right_margin) if section.right_margin else None,
        })

    core_props = doc.core_properties
    props_info = {
        "title": core_props.title,
        "author": core_props.author,
        "subject": core_props.subject,
        "created": str(core_props.created) if core_props.created else None,
        "modified": str(core_props.modified) if core_props.modified else None,
    }

    styles_available = [s.name for s in doc.styles if not s.name.startswith("_")]

    return {
        "file_path": abs_path,
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables),
        "sections_count": len(doc.sections),
        "core_properties": props_info,
        "available_styles": styles_available[:50],
        "paragraphs": paragraphs_info[:30],
        "tables": tables_info[:10],
        "sections": sections_info,
    }


def modify_paragraph_style(
    file_path: str,
    paragraph_index: int,
    style_name: str | None = None,
    alignment: str | None = None,
    font_name: str | None = None,
    font_size: float | None = None,
    font_size_unit: str = "pt",
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    font_color: str | None = None,
    line_spacing: float | None = None,
    space_before: float | None = None,
    space_after: float | None = None,
    spacing_unit: str = "pt",
    first_line_indent: float | None = None,
    first_line_indent_unit: str = "cm",
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return {"success": False, "error": f"段落索引超出范围: {paragraph_index}, 共 {len(doc.paragraphs)} 个段落"}

    para = doc.paragraphs[paragraph_index]

    if style_name:
        try:
            para.style = doc.styles[style_name]
        except KeyError:
            return {"success": False, "error": f"样式不存在: {style_name}"}

    if alignment and alignment.lower() in ALIGNMENT_MAP:
        para.alignment = ALIGNMENT_MAP[alignment.lower()]

    if line_spacing is not None:
        pf = para.paragraph_format
        pf.line_spacing = Pt(line_spacing)

    if space_before is not None:
        para.paragraph_format.space_before = _resolve_unit(space_before, spacing_unit)

    if space_after is not None:
        para.paragraph_format.space_after = _resolve_unit(space_after, spacing_unit)

    if first_line_indent is not None:
        para.paragraph_format.first_line_indent = _resolve_unit(first_line_indent, first_line_indent_unit)

    for run in para.runs:
        font = run.font
        if font_name:
            font.name = font_name
            r_element = run._element
            r_element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        if font_size is not None:
            font.size = _resolve_unit(font_size, font_size_unit)
        if bold is not None:
            font.bold = bold
        if italic is not None:
            font.italic = italic
        if underline is not None:
            font.underline = underline
        if font_color:
            try:
                font.color.rgb = RGBColor.from_string(font_color.replace("#", ""))
            except ValueError:
                return {"success": False, "error": f"无效的颜色值: {font_color}"}

    doc.save(abs_path)
    return {"success": True, "message": f"段落 {paragraph_index} 样式已更新"}


def modify_run_style(
    file_path: str,
    paragraph_index: int,
    run_index: int,
    font_name: str | None = None,
    font_size: float | None = None,
    font_size_unit: str = "pt",
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    font_color: str | None = None,
    highlight_color: str | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return {"success": False, "error": f"段落索引超出范围: {paragraph_index}"}

    para = doc.paragraphs[paragraph_index]
    if run_index < 0 or run_index >= len(para.runs):
        return {"success": False, "error": f"Run 索引超出范围: {run_index}, 共 {len(para.runs)} 个 Run"}

    run = para.runs[run_index]
    font = run.font

    if font_name:
        font.name = font_name
        r_element = run._element
        r_element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if font_size is not None:
        font.size = _resolve_unit(font_size, font_size_unit)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if underline is not None:
        font.underline = underline
    if font_color:
        try:
            font.color.rgb = RGBColor.from_string(font_color.replace("#", ""))
        except ValueError:
            return {"success": False, "error": f"无效的颜色值: {font_color}"}
    if highlight_color:
        from docx.enum.text import WD_COLOR_INDEX
        highlight_map = {
            "yellow": WD_COLOR_INDEX.YELLOW,
            "green": WD_COLOR_INDEX.GREEN,
            "cyan": WD_COLOR_INDEX.CYAN,
            "magenta": WD_COLOR_INDEX.MAGENTA,
            "blue": WD_COLOR_INDEX.BLUE,
            "red": WD_COLOR_INDEX.RED,
            "dark_blue": WD_COLOR_INDEX.DARK_BLUE,
            "dark_cyan": WD_COLOR_INDEX.DARK_CYAN,
            "dark_green": WD_COLOR_INDEX.DARK_GREEN,
            "dark_magenta": WD_COLOR_INDEX.DARK_MAGENTA,
            "dark_red": WD_COLOR_INDEX.DARK_RED,
            "dark_yellow": WD_COLOR_INDEX.DARK_YELLOW,
            "dark_gray": WD_COLOR_INDEX.DARK_GRAY,
            "light_gray": WD_COLOR_INDEX.LIGHT_GRAY,
            "black": WD_COLOR_INDEX.BLACK,
        }
        if highlight_color.lower() in highlight_map:
            font.highlight_color = highlight_map[highlight_color.lower()]
        else:
            return {"success": False, "error": f"不支持的高亮颜色: {highlight_color}"}

    doc.save(abs_path)
    return {"success": True, "message": f"段落 {paragraph_index} Run {run_index} 样式已更新"}


def replace_text(
    file_path: str,
    old_text: str,
    new_text: str,
    match_case: bool = True,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    replace_count = 0

    for para in doc.paragraphs:
        full_text = para.text
        if match_case:
            if old_text not in full_text:
                continue
        else:
            if old_text.lower() not in full_text.lower():
                continue

        if len(para.runs) == 0:
            continue

        if len(para.runs) == 1:
            run = para.runs[0]
            if match_case:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)
                    replace_count += 1
            else:
                if old_text.lower() in run.text.lower():
                    run.text = run.text.replace(old_text, new_text)
                    replace_count += 1
        else:
            if match_case:
                new_full = full_text.replace(old_text, new_text)
            else:
                import re
                new_full = re.sub(re.escape(old_text), new_text, full_text, flags=re.IGNORECASE)

            if new_full != full_text:
                first_run = para.runs[0]
                first_run.text = new_full
                for run in para.runs[1:]:
                    run.text = ""
                replace_count += 1

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    full_text = para.text
                    if match_case:
                        if old_text not in full_text:
                            continue
                    else:
                        if old_text.lower() not in full_text.lower():
                            continue

                    if len(para.runs) == 0:
                        continue

                    if match_case:
                        new_full = full_text.replace(old_text, new_text)
                    else:
                        import re
                        new_full = re.sub(re.escape(old_text), new_text, full_text, flags=re.IGNORECASE)

                    if new_full != full_text:
                        first_run = para.runs[0]
                        first_run.text = new_full
                        for run in para.runs[1:]:
                            run.text = ""
                        replace_count += 1

    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header and header.is_linked_to_previous is False:
                for para in header.paragraphs:
                    if para.text and old_text in para.text:
                        for run in para.runs:
                            if old_text in run.text:
                                run.text = run.text.replace(old_text, new_text)
                                replace_count += 1
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer and footer.is_linked_to_previous is False:
                for para in footer.paragraphs:
                    if para.text and old_text in para.text:
                        for run in para.runs:
                            if old_text in run.text:
                                run.text = run.text.replace(old_text, new_text)
                                replace_count += 1

    doc.save(abs_path)
    return {"success": True, "replace_count": replace_count, "message": f"已替换 {replace_count} 处"}


def add_paragraph(
    file_path: str,
    text: str,
    style_name: str | None = None,
    alignment: str | None = None,
    font_name: str | None = None,
    font_size: float | None = None,
    font_size_unit: str = "pt",
    bold: bool | None = None,
    italic: bool | None = None,
    font_color: str | None = None,
    insert_after_index: int | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if insert_after_index is not None:
        if insert_after_index < 0 or insert_after_index >= len(doc.paragraphs):
            return {"success": False, "error": f"段落索引超出范围: {insert_after_index}"}
        ref_para = doc.paragraphs[insert_after_index]
        new_para = doc.add_paragraph()
        ref_para._element.addnext(new_para._element)
    else:
        new_para = doc.add_paragraph()

    if style_name:
        try:
            new_para.style = doc.styles[style_name]
        except KeyError:
            return {"success": False, "error": f"样式不存在: {style_name}"}

    if alignment and alignment.lower() in ALIGNMENT_MAP:
        new_para.alignment = ALIGNMENT_MAP[alignment.lower()]

    run = new_para.add_run(text)
    font = run.font
    if font_name:
        font.name = font_name
        r_element = run._element
        r_element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if font_size is not None:
        font.size = _resolve_unit(font_size, font_size_unit)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if font_color:
        try:
            font.color.rgb = RGBColor.from_string(font_color.replace("#", ""))
        except ValueError:
            return {"success": False, "error": f"无效的颜色值: {font_color}"}

    doc.save(abs_path)
    return {"success": True, "message": f"已添加段落: {text[:50]}"}


def delete_paragraph(
    file_path: str,
    paragraph_index: int,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return {"success": False, "error": f"段落索引超出范围: {paragraph_index}"}

    para = doc.paragraphs[paragraph_index]
    p_element = para._element
    p_element.getparent().remove(p_element)

    doc.save(abs_path)
    return {"success": True, "message": f"已删除段落 {paragraph_index}"}


def modify_table_style(
    file_path: str,
    table_index: int,
    style_name: str | None = None,
    alignment: str | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if table_index < 0 or table_index >= len(doc.tables):
        return {"success": False, "error": f"表格索引超出范围: {table_index}"}

    table = doc.tables[table_index]

    if style_name:
        try:
            table.style = doc.styles[style_name]
        except KeyError:
            return {"success": False, "error": f"表格样式不存在: {style_name}"}

    if alignment and alignment.lower() in TABLE_ALIGNMENT_MAP:
        table.alignment = TABLE_ALIGNMENT_MAP[alignment.lower()]

    doc.save(abs_path)
    return {"success": True, "message": f"表格 {table_index} 样式已更新"}


def modify_table_cell(
    file_path: str,
    table_index: int,
    row_index: int,
    col_index: int,
    text: str | None = None,
    font_name: str | None = None,
    font_size: float | None = None,
    font_size_unit: str = "pt",
    bold: bool | None = None,
    italic: bool | None = None,
    font_color: str | None = None,
    alignment: str | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if table_index < 0 or table_index >= len(doc.tables):
        return {"success": False, "error": f"表格索引超出范围: {table_index}"}

    table = doc.tables[table_index]
    if row_index < 0 or row_index >= len(table.rows):
        return {"success": False, "error": f"行索引超出范围: {row_index}"}
    if col_index < 0 or col_index >= len(table.columns):
        return {"success": False, "error": f"列索引超出范围: {col_index}"}

    cell = table.rows[row_index].cells[col_index]

    if text is not None:
        if len(cell.paragraphs) > 0:
            cell.paragraphs[0].clear()
            para = cell.paragraphs[0]
        else:
            para = cell.add_paragraph()
        run = para.add_run(text)
    else:
        para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        run = para.runs[0] if para.runs else para.add_run("")

    if alignment and alignment.lower() in ALIGNMENT_MAP:
        para.alignment = ALIGNMENT_MAP[alignment.lower()]

    font = run.font
    if font_name:
        font.name = font_name
        r_element = run._element
        r_element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if font_size is not None:
        font.size = _resolve_unit(font_size, font_size_unit)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if font_color:
        try:
            font.color.rgb = RGBColor.from_string(font_color.replace("#", ""))
        except ValueError:
            return {"success": False, "error": f"无效的颜色值: {font_color}"}

    doc.save(abs_path)
    return {"success": True, "message": f"表格 {table_index} 单元格 ({row_index},{col_index}) 已更新"}


def modify_page_margins(
    file_path: str,
    section_index: int = 0,
    top: float | None = None,
    bottom: float | None = None,
    left: float | None = None,
    right: float | None = None,
    unit: str = "cm",
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if section_index < 0 or section_index >= len(doc.sections):
        return {"success": False, "error": f"节索引超出范围: {section_index}"}

    section = doc.sections[section_index]
    if top is not None:
        section.top_margin = _resolve_unit(top, unit)
    if bottom is not None:
        section.bottom_margin = _resolve_unit(bottom, unit)
    if left is not None:
        section.left_margin = _resolve_unit(left, unit)
    if right is not None:
        section.right_margin = _resolve_unit(right, unit)

    doc.save(abs_path)
    return {"success": True, "message": f"节 {section_index} 页边距已更新"}


def modify_page_size(
    file_path: str,
    section_index: int = 0,
    width: float | None = None,
    height: float | None = None,
    unit: str = "cm",
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if section_index < 0 or section_index >= len(doc.sections):
        return {"success": False, "error": f"节索引超出范围: {section_index}"}

    section = doc.sections[section_index]
    if width is not None:
        section.page_width = _resolve_unit(width, unit)
    if height is not None:
        section.page_height = _resolve_unit(height, unit)

    doc.save(abs_path)
    return {"success": True, "message": f"节 {section_index} 页面尺寸已更新"}


def modify_header_footer(
    file_path: str,
    section_index: int = 0,
    header_text: str | None = None,
    footer_text: str | None = None,
    header_font_name: str | None = None,
    header_font_size: float | None = None,
    header_font_size_unit: str = "pt",
    footer_font_name: str | None = None,
    footer_font_size: float | None = None,
    footer_font_size_unit: str = "pt",
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if section_index < 0 or section_index >= len(doc.sections):
        return {"success": False, "error": f"节索引超出范围: {section_index}"}

    section = doc.sections[section_index]
    changes = []

    if header_text is not None:
        header = section.header
        header.is_linked_to_previous = False
        if len(header.paragraphs) > 0:
            para = header.paragraphs[0]
            para.clear()
        else:
            para = header.add_paragraph()
        run = para.add_run(header_text)
        if header_font_name:
            run.font.name = header_font_name
            run._element.rPr.rFonts.set(qn("w:eastAsia"), header_font_name)
        if header_font_size is not None:
            run.font.size = _resolve_unit(header_font_size, header_font_size_unit)
        changes.append("页眉")

    if footer_text is not None:
        footer = section.footer
        footer.is_linked_to_previous = False
        if len(footer.paragraphs) > 0:
            para = footer.paragraphs[0]
            para.clear()
        else:
            para = footer.add_paragraph()
        run = para.add_run(footer_text)
        if footer_font_name:
            run.font.name = footer_font_name
            run._element.rPr.rFonts.set(qn("w:eastAsia"), footer_font_name)
        if footer_font_size is not None:
            run.font.size = _resolve_unit(footer_font_size, footer_font_size_unit)
        changes.append("页脚")

    doc.save(abs_path)
    return {"success": True, "message": f"节 {section_index} {'、'.join(changes)}已更新"}


def batch_modify_paragraphs(
    file_path: str,
    style_name: str | None = None,
    alignment: str | None = None,
    font_name: str | None = None,
    font_size: float | None = None,
    font_size_unit: str = "pt",
    bold: bool | None = None,
    italic: bool | None = None,
    font_color: str | None = None,
    line_spacing: float | None = None,
    space_before: float | None = None,
    space_after: float | None = None,
    spacing_unit: str = "pt",
    first_line_indent: float | None = None,
    first_line_indent_unit: str = "cm",
    paragraph_indices: list[int] | None = None,
    heading_only: bool = False,
    body_only: bool = False,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    target_indices = []
    if paragraph_indices is not None:
        target_indices = paragraph_indices
    else:
        for i, p in enumerate(doc.paragraphs):
            style = p.style.name if p.style else ""
            if heading_only and not style.startswith("Heading"):
                continue
            if body_only and style.startswith("Heading"):
                continue
            target_indices.append(i)

    modified_count = 0
    for idx in target_indices:
        if idx < 0 or idx >= len(doc.paragraphs):
            continue
        para = doc.paragraphs[idx]

        if style_name:
            try:
                para.style = doc.styles[style_name]
            except KeyError:
                continue

        if alignment and alignment.lower() in ALIGNMENT_MAP:
            para.alignment = ALIGNMENT_MAP[alignment.lower()]

        if line_spacing is not None:
            para.paragraph_format.line_spacing = Pt(line_spacing)
        if space_before is not None:
            para.paragraph_format.space_before = _resolve_unit(space_before, spacing_unit)
        if space_after is not None:
            para.paragraph_format.space_after = _resolve_unit(space_after, spacing_unit)
        if first_line_indent is not None:
            para.paragraph_format.first_line_indent = _resolve_unit(first_line_indent, first_line_indent_unit)

        for run in para.runs:
            font = run.font
            if font_name:
                font.name = font_name
                run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
            if font_size is not None:
                font.size = _resolve_unit(font_size, font_size_unit)
            if bold is not None:
                font.bold = bold
            if italic is not None:
                font.italic = italic
            if font_color:
                try:
                    font.color.rgb = RGBColor.from_string(font_color.replace("#", ""))
                except ValueError:
                    pass

        modified_count += 1

    doc.save(abs_path)
    return {"success": True, "modified_count": modified_count, "message": f"已批量修改 {modified_count} 个段落"}


def modify_core_properties(
    file_path: str,
    title: str | None = None,
    author: str | None = None,
    subject: str | None = None,
    keywords: str | None = None,
    comments: str | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    props = doc.core_properties
    if title is not None:
        props.title = title
    if author is not None:
        props.author = author
    if subject is not None:
        props.subject = subject
    if keywords is not None:
        props.keywords = keywords
    if comments is not None:
        props.comments = comments

    doc.save(abs_path)
    return {"success": True, "message": "文档属性已更新"}


def add_table(
    file_path: str,
    rows: int,
    cols: int,
    data: list[list[str]] | None = None,
    style_name: str | None = None,
    header_row: bool = True,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    table = doc.add_table(rows=rows, cols=cols)

    if style_name:
        try:
            table.style = doc.styles[style_name]
        except KeyError:
            pass

    if data:
        for i, row_data in enumerate(data):
            if i >= rows:
                break
            for j, cell_text in enumerate(row_data):
                if j >= cols:
                    break
                cell = table.rows[i].cells[j]
                cell.text = str(cell_text)

    doc.save(abs_path)
    return {"success": True, "message": f"已添加 {rows}x{cols} 表格"}


def add_page_break(
    file_path: str,
    insert_after_index: int | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    if insert_after_index is not None:
        if insert_after_index < 0 or insert_after_index >= len(doc.paragraphs):
            return {"success": False, "error": f"段落索引超出范围: {insert_after_index}"}
        para = doc.paragraphs[insert_after_index]
        run = para.add_run()
        run.add_break()
    else:
        doc.add_page_break()

    doc.save(abs_path)
    return {"success": True, "message": "已添加分页符"}


def add_heading(
    file_path: str,
    text: str,
    level: int = 1,
    font_name: str | None = None,
    font_size: float | None = None,
    font_size_unit: str = "pt",
    font_color: str | None = None,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    heading = doc.add_heading(text, level=level)

    for run in heading.runs:
        font = run.font
        if font_name:
            font.name = font_name
            run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        if font_size is not None:
            font.size = _resolve_unit(font_size, font_size_unit)
        if font_color:
            try:
                font.color.rgb = RGBColor.from_string(font_color.replace("#", ""))
            except ValueError:
                pass

    doc.save(abs_path)
    return {"success": True, "message": f"已添加标题 (级别 {level}): {text[:50]}"}


def insert_image(
    file_path: str,
    image_path: str,
    width: float | None = None,
    height: float | None = None,
    unit: str = "cm",
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    image_abs = os.path.abspath(image_path)

    if not os.path.isfile(image_abs):
        return {"success": False, "error": f"图片文件不存在: {image_abs}"}

    doc = Document(abs_path)

    kwargs = {}
    if width is not None:
        kwargs["width"] = _resolve_unit(width, unit)
    if height is not None:
        kwargs["height"] = _resolve_unit(height, unit)

    doc.add_picture(image_abs, **kwargs)

    doc.save(abs_path)
    return {"success": True, "message": f"已插入图片: {image_abs}"}


def find_text(
    file_path: str,
    search_text: str,
    match_case: bool = True,
) -> dict[str, Any]:
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)

    results = []

    for i, para in enumerate(doc.paragraphs):
        text = para.text
        found = (search_text in text) if match_case else (search_text.lower() in text.lower())
        if found:
            results.append({
                "type": "paragraph",
                "index": i,
                "text": text[:200],
                "style": para.style.name if para.style else None,
            })

    for i, table in enumerate(doc.tables):
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                text = cell.text
                found = (search_text in text) if match_case else (search_text.lower() in text.lower())
                if found:
                    results.append({
                        "type": "table_cell",
                        "table_index": i,
                        "row": r,
                        "col": c,
                        "text": text[:200],
                    })

    return {
        "success": True,
        "search_text": search_text,
        "found_count": len(results),
        "results": results[:50],
    }
