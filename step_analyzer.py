#!/root/.pyenv/versions/3.11.15/bin/python3
# -*- coding: utf-8 -*-
"""STEP文件分析工具 - 解析STEP文件并生成Word分析报告"""

import argparse
import os
import sys
import math

from OCP.STEPControl import STEPControl_Reader
from OCP.TopExp import TopExp_Explorer
from OCP.TopAbs import (
    TopAbs_VERTEX, TopAbs_EDGE, TopAbs_WIRE, TopAbs_FACE,
    TopAbs_SHELL, TopAbs_SOLID
)
from OCP.TopoDS import TopoDS, TopoDS_Compound
from OCP.TopExp import TopExp
from OCP.BRepAdaptor import BRepAdaptor_Surface, BRepAdaptor_Curve
from OCP.GeomAbs import (
    GeomAbs_Plane, GeomAbs_Cylinder, GeomAbs_Cone, GeomAbs_Sphere,
    GeomAbs_Torus, GeomAbs_BSplineSurface, GeomAbs_BezierSurface,
    GeomAbs_SurfaceOfRevolution, GeomAbs_SurfaceOfExtrusion, GeomAbs_OtherSurface,
    GeomAbs_Line, GeomAbs_Circle, GeomAbs_Ellipse, GeomAbs_Hyperbola,
    GeomAbs_Parabola, GeomAbs_BezierCurve, GeomAbs_BSplineCurve, GeomAbs_OtherCurve
)
from OCP.Bnd import Bnd_Box
from OCP.BRepBndLib import BRepBndLib
from OCP.BRepGProp import BRepGProp
from OCP.GProp import GProp_GProps
from OCP.BRepExtrema import BRepExtrema_DistShapeShape
from OCP.gp import gp_Vec
from OCP.TopTools import TopTools_IndexedMapOfShape

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# ============================================================
# 几何类型名称映射
# ============================================================
FACE_TYPE_NAMES = {
    GeomAbs_Plane: "平面",
    GeomAbs_Cylinder: "圆柱面",
    GeomAbs_Cone: "圆锥面",
    GeomAbs_Sphere: "球面",
    GeomAbs_Torus: "环面",
    GeomAbs_BSplineSurface: "B样条曲面",
    GeomAbs_BezierSurface: "贝塞尔曲面",
    GeomAbs_SurfaceOfRevolution: "旋转曲面",
    GeomAbs_SurfaceOfExtrusion: "拉伸曲面",
    GeomAbs_OtherSurface: "其他曲面",
}

EDGE_TYPE_NAMES = {
    GeomAbs_Line: "直线",
    GeomAbs_Circle: "圆",
    GeomAbs_Ellipse: "椭圆",
    GeomAbs_Hyperbola: "双曲线",
    GeomAbs_Parabola: "抛物线",
    GeomAbs_BezierCurve: "贝塞尔曲线",
    GeomAbs_BSplineCurve: "B样条曲线",
    GeomAbs_OtherCurve: "其他曲线",
}


# ============================================================
# STEP文件读取
# ============================================================
def read_step_file(filepath):
    """读取STEP文件，返回TopoDS_Shape"""
    reader = STEPControl_Reader()
    status = reader.ReadFile(filepath)
    if status != 1:  # IFSelect_RetDone
        return None
    reader.TransferRoots()
    return reader.OneShape()


# ============================================================
# 拓扑信息统计
# ============================================================
def count_topology(shape):
    """统计顶点、边、线、面、壳、实体数量"""
    counts = {}
    for name, topo_type in [
        ("顶点", TopAbs_VERTEX),
        ("边", TopAbs_EDGE),
        ("线", TopAbs_WIRE),
        ("面", TopAbs_FACE),
        ("壳", TopAbs_SHELL),
        ("实体", TopAbs_SOLID),
    ]:
        exp = TopExp_Explorer(shape, topo_type)
        count = 0
        while exp.More():
            count += 1
            exp.Next()
        counts[name] = count
    return counts


# ============================================================
# 几何类型统计
# ============================================================
def count_face_types(shape):
    """统计各面类型数量"""
    type_counts = {}
    exp = TopExp_Explorer(shape, TopAbs_FACE)
    while exp.More():
        face = TopoDS.Face_s(exp.Current())
        adaptor = BRepAdaptor_Surface(face)
        t = adaptor.GetType()
        name = FACE_TYPE_NAMES.get(t, f"未知({t})")
        type_counts[name] = type_counts.get(name, 0) + 1
        exp.Next()
    return type_counts


def count_edge_types(shape):
    """统计各边类型数量"""
    type_counts = {}
    exp = TopExp_Explorer(shape, TopAbs_EDGE)
    while exp.More():
        edge = TopoDS.Edge_s(exp.Current())
        adaptor = BRepAdaptor_Curve(edge)
        t = adaptor.GetType()
        name = EDGE_TYPE_NAMES.get(t, f"未知({t})")
        type_counts[name] = type_counts.get(name, 0) + 1
        exp.Next()
    return type_counts


# ============================================================
# 尺寸信息
# ============================================================
def get_bounding_box(shape):
    """获取包围盒"""
    bbox = Bnd_Box()
    BRepBndLib.Add_s(shape, bbox)
    xmin, ymin, zmin, xmax, ymax, zmax = bbox.Get()
    return {
        "x_min": xmin, "x_max": xmax, "x_size": xmax - xmin,
        "y_min": ymin, "y_max": ymax, "y_size": ymax - ymin,
        "z_min": zmin, "z_max": zmax, "z_size": zmax - zmin,
    }


def get_volume(shape):
    """获取体积"""
    props = GProp_GProps()
    BRepGProp.VolumeProperties_s(shape, props)
    return props.Mass(), props.CentreOfMass()


def get_surface_area(shape):
    """获取表面积"""
    props = GProp_GProps()
    BRepGProp.SurfaceProperties_s(shape, props)
    return props.Mass()


# ============================================================
# 孔特征分析
# ============================================================
def analyze_holes(shape):
    """分析孔特征"""
    holes = []
    exp = TopExp_Explorer(shape, TopAbs_FACE)
    while exp.More():
        face = TopoDS.Face_s(exp.Current())
        adaptor = BRepAdaptor_Surface(face)
        surf_type = adaptor.GetType()

        # 统计面内的线数量
        wire_exp = TopExp_Explorer(face, TopAbs_WIRE)
        wire_count = 0
        while wire_exp.More():
            wire_count += 1
            wire_exp.Next()

        # 孔面条件：圆柱面/圆锥面且有内线(wire_count > 1)，或平面有内线
        is_hole_face = False
        if surf_type in (GeomAbs_Cylinder, GeomAbs_Cone) and wire_count > 1:
            is_hole_face = True
        elif surf_type == GeomAbs_Plane and wire_count > 1:
            is_hole_face = True

        if not is_hole_face:
            exp.Next()
            continue

        # 分析每条内线的边类型
        wire_exp = TopExp_Explorer(face, TopAbs_WIRE)
        wire_idx = 0
        inner_wires_data = []  # [(edge_types_list, [edge_adaptor_list]), ...]
        while wire_exp.More():
            if wire_idx > 0:  # 内线
                wire = wire_exp.Current()
                edge_exp = TopExp_Explorer(wire, TopAbs_EDGE)
                edge_types = []
                edge_adaptors = []
                while edge_exp.More():
                    edge = TopoDS.Edge_s(edge_exp.Current())
                    curve_adaptor = BRepAdaptor_Curve(edge)
                    edge_types.append(curve_adaptor.GetType())
                    edge_adaptors.append(curve_adaptor)
                    edge_exp.Next()
                inner_wires_data.append((edge_types, edge_adaptors))
            wire_idx += 1
            wire_exp.Next()

        # 提取参数
        if surf_type in (GeomAbs_Cylinder, GeomAbs_Cone):
            # 圆柱/圆锥面孔面 - 整个面就是一个孔
            hole_info = {"surface_type": surf_type}
            if surf_type == GeomAbs_Cylinder:
                cyl = adaptor.Cylinder()
                hole_info["radius"] = cyl.Radius()
                hole_info["diameter"] = cyl.Radius() * 2
                loc = cyl.Location()
                axis = cyl.Axis()
                d = axis.Direction()
                hole_info["axis_location"] = (loc.X(), loc.Y(), loc.Z())
                hole_info["axis_direction"] = (d.X(), d.Y(), d.Z())
            elif surf_type == GeomAbs_Cone:
                cone = adaptor.Cone()
                hole_info["radius"] = cone.RefRadius()
                hole_info["diameter"] = cone.RefRadius() * 2
                loc = cone.Location()
                axis = cone.Axis()
                d = axis.Direction()
                hole_info["axis_location"] = (loc.X(), loc.Y(), loc.Z())
                hole_info["axis_direction"] = (d.X(), d.Y(), d.Z())
                hole_info["semi_angle"] = cone.SemiAngle()

            # 分类
            all_edge_types = []
            for edge_types, _ in inner_wires_data:
                all_edge_types.extend(edge_types)
            hole_info["classification"] = _classify_hole(
                inner_wires_data, all_edge_types, surf_type
            )
            holes.append(hole_info)

        elif surf_type == GeomAbs_Plane:
            # 平面上的孔 - 每条内线代表一个孔
            for edge_types, edge_adaptors in inner_wires_data:
                hole_info = {"surface_type": surf_type}

                # 从内线的圆边获取半径和轴信息
                found_circle = False
                for adaptor_c in edge_adaptors:
                    if adaptor_c.GetType() == GeomAbs_Circle:
                        circ = adaptor_c.Circle()
                        hole_info["radius"] = circ.Radius()
                        hole_info["diameter"] = circ.Radius() * 2
                        loc = circ.Location()
                        d = circ.Axis().Direction()
                        hole_info["axis_location"] = (loc.X(), loc.Y(), loc.Z())
                        hole_info["axis_direction"] = (d.X(), d.Y(), d.Z())
                        found_circle = True
                        break

                # 纯直线内Wire（矩形/方形槽）：计算长宽
                if not found_circle and all(t == GeomAbs_Line for t in edge_types):
                    # 获取每条边的长度
                    from OCP.BRepGProp import BRepGProp as BRepGProp2
                    from OCP.GProp import GProp_GProps as GProp_GProps2
                    edge_lengths = []
                    for adaptor_c in edge_adaptors:
                        # 获取边的首尾点距离
                        first_param = adaptor_c.FirstParameter()
                        last_param = adaptor_c.LastParameter()
                        p1 = adaptor_c.Value(first_param)
                        p2 = adaptor_c.Value(last_param)
                        length = math.sqrt(
                            (p2.X()-p1.X())**2 + (p2.Y()-p1.Y())**2 + (p2.Z()-p1.Z())**2
                        )
                        edge_lengths.append(length)
                    # 去重排序，取两个不同的长度作为长和宽
                    unique_lengths = sorted(set(round(l, 3) for l in edge_lengths))
                    if len(unique_lengths) >= 2:
                        hole_info["slot_length"] = max(unique_lengths)
                        hole_info["slot_width"] = min(unique_lengths)
                    elif len(unique_lengths) == 1:
                        hole_info["slot_length"] = unique_lengths[0]
                        hole_info["slot_width"] = unique_lengths[0]
                    else:
                        hole_info["slot_length"] = 0
                        hole_info["slot_width"] = 0

                # 分类
                hole_info["classification"] = _classify_hole(
                    [(edge_types, edge_adaptors)], edge_types, surf_type
                )
                holes.append(hole_info)

        exp.Next()

    return holes


def _classify_hole(inner_wires_data, all_edge_types, surf_type):
    """分类孔类型"""
    has_bspline = GeomAbs_BSplineCurve in all_edge_types
    has_circle = GeomAbs_Circle in all_edge_types
    has_line = GeomAbs_Line in all_edge_types
    has_other = any(
        t not in (GeomAbs_Circle, GeomAbs_Line)
        for t in all_edge_types
    )

    # 检查周围面是否为B样条曲面
    surrounding_bspline = (surf_type == GeomAbs_BSplineSurface)

    if has_bspline or surrounding_bspline:
        return "不规则孔（含B样条曲线）"
    elif has_other:
        return "不规则孔"
    elif len(inner_wires_data) == 1 and len(inner_wires_data[0][0]) == 1 and inner_wires_data[0][0][0] == GeomAbs_Circle:
        return "规则孔（单圆）"
    elif has_circle and not has_line:
        return "规则孔（多段圆弧）"
    elif has_circle and has_line:
        return "规则孔（直线+圆弧，如键槽）"
    elif has_line and not has_circle:
        # 纯直线组成的内轮廓（矩形/方形槽）
        return "规则孔（矩形/方形槽）"
    else:
        return "不规则孔"


# ============================================================
# 管状/方通结构分析
# ============================================================
def _vec_almost_equal(v1, v2, tol=1e-3):
    """比较两个方向向量是否近似相等（考虑反方向）"""
    dot = abs(v1[0]*v2[0] + v1[1]*v2[1] + v1[2]*v2[2])
    return dot > (1.0 - tol)


def _axis_perp_position(axis_dir, axis_loc):
    """获取轴线上垂直于轴方向的位置（用于判断同轴）"""
    # 将位置投影到垂直于轴方向的平面上
    # 实际上用轴位置本身，但只取垂直于轴方向的分量
    d = axis_dir
    # 投影：loc - (loc·d)*d
    dot = axis_loc[0]*d[0] + axis_loc[1]*d[1] + axis_loc[2]*d[2]
    perp = (axis_loc[0] - dot*d[0], axis_loc[1] - dot*d[1], axis_loc[2] - dot*d[2])
    return perp


def _perp_distance(p1, p2):
    """两点间距离"""
    return math.sqrt((p1[0]-p2[0])**2 + (p1[1]-p2[1])**2 + (p1[2]-p2[2])**2)


def analyze_tube_structures(shape):
    """分析管状结构（同轴圆柱面对）"""
    # 收集所有圆柱面信息
    cyl_faces = []
    exp = TopExp_Explorer(shape, TopAbs_FACE)
    while exp.More():
        face = TopoDS.Face_s(exp.Current())
        adaptor = BRepAdaptor_Surface(face)
        if adaptor.GetType() == GeomAbs_Cylinder:
            cyl = adaptor.Cylinder()
            loc = cyl.Location()
            axis = cyl.Axis()
            d = axis.Direction()
            radius = cyl.Radius()
            cyl_faces.append({
                "face": face,
                "radius": radius,
                "axis_location": (loc.X(), loc.Y(), loc.Z()),
                "axis_direction": (d.X(), d.Y(), d.Z()),
                "perp_pos": _axis_perp_position((d.X(), d.Y(), d.Z()), (loc.X(), loc.Y(), loc.Z())),
            })
        exp.Next()

    # 找同轴圆柱面对（相同轴方向，相同垂直位置）
    tubes = []
    used = set()
    for i, c1 in enumerate(cyl_faces):
        if i in used:
            continue
        for j, c2 in enumerate(cyl_faces):
            if j <= i or j in used:
                continue
            # 检查轴方向是否相同
            if not _vec_almost_equal(c1["axis_direction"], c2["axis_direction"]):
                continue
            # 检查垂直位置是否相同（同轴）
            if _perp_distance(c1["perp_pos"], c2["perp_pos"]) > 1e-3:
                continue
            # 找到同轴对
            outer = c1 if c1["radius"] > c2["radius"] else c2
            inner = c2 if c1["radius"] > c2["radius"] else c1
            wall_thickness = outer["radius"] - inner["radius"]
            if wall_thickness < 1e-6:
                continue
            tubes.append({
                "outer_radius": outer["radius"],
                "inner_radius": inner["radius"],
                "wall_thickness": wall_thickness,
                "axis_direction": c1["axis_direction"],
                "axis_location": c1["axis_location"],
                "outer_face": outer["face"],
                "inner_face": inner["face"],
            })
            used.add(i)
            used.add(j)

    # 计算管长度（投影包围盒到轴方向）
    for tube in tubes:
        bbox = Bnd_Box()
        BRepBndLib.Add_s(tube["outer_face"], bbox)
        xmin, ymin, zmin, xmax, ymax, zmax = bbox.Get()
        d = tube["axis_direction"]
        # 投影包围盒对角线到轴方向
        proj_min = xmin*d[0] + ymin*d[1] + zmin*d[2]
        proj_max = xmax*d[0] + ymax*d[1] + zmax*d[2]
        tube["length"] = abs(proj_max - proj_min)
        tube["outer_diameter"] = tube["outer_radius"] * 2
        tube["inner_diameter"] = tube["inner_radius"] * 2

    return tubes


def analyze_square_tube_structures(shape):
    """分析方通结构（平行平面对）"""
    # 收集所有平面信息
    plane_faces = []
    exp = TopExp_Explorer(shape, TopAbs_FACE)
    while exp.More():
        face = TopoDS.Face_s(exp.Current())
        adaptor = BRepAdaptor_Surface(face)
        if adaptor.GetType() == GeomAbs_Plane:
            pln = adaptor.Plane()
            pos = pln.Position()
            loc = pos.Location()
            d = pos.Direction()
            plane_faces.append({
                "face": face,
                "location": (loc.X(), loc.Y(), loc.Z()),
                "direction": (d.X(), d.Y(), d.Z()),
            })
        exp.Next()

    # 找平行平面对（方向相同或相反，位置不同）
    square_tubes = []
    used = set()
    for i, p1 in enumerate(plane_faces):
        if i in used:
            continue
        for j, p2 in enumerate(plane_faces):
            if j <= i or j in used:
                continue
            d1 = p1["direction"]
            d2 = p2["direction"]
            # 检查是否平行（方向相同或相反）
            dot = d1[0]*d2[0] + d1[1]*d2[1] + d1[2]*d2[2]
            if abs(abs(dot) - 1.0) > 1e-3:
                continue
            # 计算平面间距
            # 使用法向量和位置差
            loc1 = p1["location"]
            loc2 = p2["location"]
            diff = (loc2[0]-loc1[0], loc2[1]-loc1[1], loc2[2]-loc1[2])
            # 间距 = diff · d1 (投影到法向量方向)
            distance = abs(diff[0]*d1[0] + diff[1]*d1[1] + diff[2]*d1[2])
            if distance < 1e-3:
                continue

            # 检查是否方向相反（内/外面）
            is_opposite = dot < 0

            square_tubes.append({
                "face1_idx": i,
                "face2_idx": j,
                "direction": d1,
                "distance": distance,
                "is_opposite_normal": is_opposite,
                "face1": p1["face"],
                "face2": p2["face"],
            })
            used.add(i)
            used.add(j)

    # 按方向分组，找方通截面
    # 方通通常有两组平行平面对（X方向和Y方向各一对）
    direction_groups = {}
    for st in square_tubes:
        d = st["direction"]
        # 归一化方向（取正方向）
        key = None
        for existing_key in direction_groups:
            if _vec_almost_equal(d, existing_key):
                key = existing_key
                break
        if key is None:
            key = d
        if key not in direction_groups:
            direction_groups[key] = []
        direction_groups[key].append(st)

    # 找有两组正交方向的方通
    result = []
    dir_list = list(direction_groups.keys())
    for i in range(len(dir_list)):
        for j in range(i+1, len(dir_list)):
            d1 = dir_list[i]
            d2 = dir_list[j]
            # 检查是否正交
            dot = abs(d1[0]*d2[0] + d1[1]*d2[1] + d1[2]*d2[2])
            if dot > 1e-3:
                continue
            group1 = direction_groups[dir_list[i]]
            group2 = direction_groups[dir_list[j]]
            # 取每组中距离最大的对（外表面）
            if group1 and group2:
                max_dist1 = max(g["distance"] for g in group1)
                max_dist2 = max(g["distance"] for g in group2)
                result.append({
                    "cross_section_dim1": max_dist1,
                    "cross_section_dim2": max_dist2,
                    "direction1": dir_list[i],
                    "direction2": dir_list[j],
                    "pairs_count_1": len(group1),
                    "pairs_count_2": len(group2),
                })

    return square_tubes, result


# ============================================================
# 壁厚分析
# ============================================================
def analyze_wall_thickness(shape, tubes):
    """分析壁厚"""
    wall_info = []

    # 从管状结构获取壁厚
    for tube in tubes:
        wall_info.append({
            "type": "管状结构",
            "wall_thickness": tube["wall_thickness"],
            "outer_radius": tube["outer_radius"],
            "inner_radius": tube["inner_radius"],
            "axis_direction": tube["axis_direction"],
        })

    # 使用BRepExtrema计算内外面之间的距离（通用方法）
    # 收集所有面
    all_faces = []
    exp = TopExp_Explorer(shape, TopAbs_FACE)
    while exp.More():
        all_faces.append(TopoDS.Face_s(exp.Current()))
        exp.Next()

    # 对管状结构的内外面计算距离
    for tube in tubes:
        try:
            dist_calc = BRepExtrema_DistShapeShape()
            dist_calc.LoadS1(tube["inner_face"])
            dist_calc.LoadS2(tube["outer_face"])
            dist_calc.Perform()
            if dist_calc.IsDone():
                min_dist = dist_calc.Value()
                wall_info.append({
                    "type": "内外面距离",
                    "wall_thickness": min_dist,
                    "axis_direction": tube["axis_direction"],
                })
        except Exception:
            pass

    return wall_info


# ============================================================
# 模型类型自动检测
# ============================================================
def detect_model_type(tubes, square_tube_pairs, square_tube_cross):
    """自动检测模型类型"""
    if square_tube_cross or square_tube_pairs:
        return "方通结构"
    if tubes:
        return "管状结构"
    return "一般机械件"


# ============================================================
# 形状结构判断与装配件分析
# ============================================================
def classify_shape_structure(shape):
    """判断STEP文件结构类型

    Returns:
        ("single_solid", solid_count) - 单零件
        ("assembly", solid_count)     - 装配件（多Solid Compound）
    """
    solid_count = 0
    exp = TopExp_Explorer(shape, TopAbs_SOLID)
    while exp.More():
        solid_count += 1
        exp.Next()

    if solid_count <= 1:
        return "single_solid", solid_count
    else:
        return "assembly", solid_count


def decompose_compound(shape):
    """将Compound分解为独立的Solid列表

    Returns:
        list of TopoDS_Solid
    """
    solids = []
    exp = TopExp_Explorer(shape, TopAbs_SOLID)
    while exp.More():
        solid = TopoDS.Solid_s(exp.Current())
        solids.append(solid)
        exp.Next()
    return solids


def _get_all_faces(shape):
    """获取形状的所有面"""
    faces = []
    exp = TopExp_Explorer(shape, TopAbs_FACE)
    while exp.More():
        faces.append(TopoDS.Face_s(exp.Current()))
        exp.Next()
    return faces


def _find_end_faces(solid):
    """识别管/方通的端面

    端面定义：沿主轴方向位于包围盒极值位置的面
    对于圆柱面：轴向端面（圆形平面）
    对于方通：轴向端面（矩形平面）

    Returns:
        list of {"face": TopoDS_Face, "position": "min"|"max", "axis": (dx,dy,dz)}
    """
    bbox = Bnd_Box()
    BRepBndLib.Add_s(solid, bbox)
    xmin, ymin, zmin, xmax, ymax, zmax = bbox.Get()

    # 确定主轴方向（最长维度）
    dims = [xmax - xmin, ymax - ymin, zmax - zmin]
    main_axis_idx = dims.index(max(dims))
    axis_dirs = [(1, 0, 0), (0, 1, 0), (0, 0, 1)]
    main_axis = axis_dirs[main_axis_idx]

    end_faces = []
    exp = TopExp_Explorer(solid, TopAbs_FACE)
    while exp.More():
        face = TopoDS.Face_s(exp.Current())
        adaptor = BRepAdaptor_Surface(face)
        surf_type = adaptor.GetType()

        # 端面通常是平面
        if surf_type != GeomAbs_Plane:
            exp.Next()
            continue

        pln = adaptor.Plane()
        pos = pln.Position()
        loc = pos.Location()
        d = pos.Direction()

        # 检查平面法向量是否与主轴平行
        dot = abs(d.X() * main_axis[0] + d.Y() * main_axis[1] + d.Z() * main_axis[2])
        if dot < 0.99:
            exp.Next()
            continue

        # 判断端面位置（min端还是max端）
        # 用面的位置在主轴方向的投影
        proj = loc.X() * main_axis[0] + loc.Y() * main_axis[1] + loc.Z() * main_axis[2]
        proj_min = xmin * main_axis[0] + ymin * main_axis[1] + zmin * main_axis[2]
        proj_max = xmax * main_axis[0] + ymax * main_axis[1] + zmax * main_axis[2]

        tolerance = max(dims) * 0.05  # 5%容差
        if abs(proj - proj_min) < tolerance:
            position = "min"
        elif abs(proj - proj_max) < tolerance:
            position = "max"
        else:
            exp.Next()
            continue

        end_faces.append({
            "face": face,
            "position": position,
            "axis": main_axis,
            "plane_loc": (loc.X(), loc.Y(), loc.Z()),
        })
        exp.Next()

    return end_faces


def analyze_assembly_connectivity(solids):
    """分析装配件中各零件的连接关系

    判断每个零件的端面是自由端（切口）还是焊接端（与其他零件连接）

    Args:
        solids: TopoDS_Solid 列表

    Returns:
        list of {
            "part_index": int,
            "model_type": str,
            "end_faces": [
                {"position": "min"|"max", "is_free": bool, "connected_to": int|None}
            ]
        }
    """
    # 先分析每个零件的基本信息和端面
    parts_info = []
    for idx, solid in enumerate(solids):
        tubes = analyze_tube_structures(solid)
        sq_pairs, sq_cross = analyze_square_tube_structures(solid)
        model_type = detect_model_type(tubes, sq_pairs, sq_cross)
        end_faces = _find_end_faces(solid)

        parts_info.append({
            "part_index": idx,
            "solid": solid,
            "model_type": model_type,
            "tubes": tubes,
            "sq_pairs": sq_pairs,
            "sq_cross": sq_cross,
            "end_faces": end_faces,
        })

    # 分析端面连接性：检查每个零件的端面是否与其他零件的面距离≈0
    for i, part in enumerate(parts_info):
        for ef in part["end_faces"]:
            ef["is_free"] = True
            ef["connected_to"] = None

            # 获取端面的包围盒
            ef_bbox = Bnd_Box()
            BRepBndLib.Add_s(ef["face"], ef_bbox)
            ef_xmin, ef_ymin, ef_zmin, ef_xmax, ef_ymax, ef_zmax = ef_bbox.Get()

            for j, other_part in enumerate(parts_info):
                if i == j:
                    continue

                # 快速排除：包围盒不相交的零件
                other_bbox = Bnd_Box()
                BRepBndLib.Add_s(other_part["solid"], other_bbox)
                oxmin, oymin, ozmin, oxmax, oymax, ozmax = other_bbox.Get()

                # 包围盒扩展一点容差后检查重叠
                tol = 1.0  # mm
                if (ef_xmax + tol < oxmin or ef_xmin - tol > oxmax or
                    ef_ymax + tol < oymin or ef_ymin - tol > oymax or
                    ef_zmax + tol < ozmin or ef_zmin - tol > ozmax):
                    continue

                # 精确距离检查
                try:
                    dist_calc = BRepExtrema_DistShapeShape()
                    dist_calc.LoadS1(ef["face"])
                    dist_calc.LoadS2(other_part["solid"])
                    dist_calc.Perform()
                    if dist_calc.IsDone() and dist_calc.Value() < 0.1:  # 0.1mm阈值
                        ef["is_free"] = False
                        ef["connected_to"] = j
                        break
                except Exception:
                    pass

    # 清理临时solid引用（不需要序列化）
    result = []
    for part in parts_info:
        result.append({
            "part_index": part["part_index"],
            "model_type": part["model_type"],
            "tubes": part["tubes"],
            "sq_pairs": part["sq_pairs"],
            "sq_cross": part["sq_cross"],
            "end_faces": [
                {
                    "position": ef["position"],
                    "is_free": ef["is_free"],
                    "connected_to": ef["connected_to"],
                }
                for ef in part["end_faces"]
            ],
        })

    return result


# ============================================================
# Word报告生成
# ============================================================
def fmt_length(val):
    """格式化长度值，3位小数+mm"""
    return f"{val:.3f} mm"


def fmt_area(val):
    """格式化面积值，2位小数+mm²"""
    return f"{val:.2f} mm²"


def fmt_volume(val):
    """格式化体积值，2位小数+mm³"""
    return f"{val:.2f} mm³"


def add_table_row(table, cells_data):
    """向表格添加一行"""
    row = table.add_row()
    for i, text in enumerate(cells_data):
        row.cells[i].text = str(text)
    return row


def _analyze_single_solid(solid):
    """分析单个Solid，返回所有分析结果"""
    topology = count_topology(solid)
    face_types = count_face_types(solid)
    edge_types = count_edge_types(solid)
    bbox = get_bounding_box(solid)
    volume, center_of_mass = get_volume(solid)
    surface_area = get_surface_area(solid)
    holes = analyze_holes(solid)
    tubes = analyze_tube_structures(solid)
    sq_pairs, sq_cross = analyze_square_tube_structures(solid)
    wall_info = analyze_wall_thickness(solid, tubes)
    model_type = detect_model_type(tubes, sq_pairs, sq_cross)

    return {
        "topology": topology,
        "face_types": face_types,
        "edge_types": edge_types,
        "bbox": bbox,
        "volume": volume,
        "center_of_mass": center_of_mass,
        "surface_area": surface_area,
        "holes": holes,
        "tubes": tubes,
        "sq_pairs": sq_pairs,
        "sq_cross": sq_cross,
        "wall_info": wall_info,
        "model_type": model_type,
    }


def _add_single_part_sections(doc, analysis, part_name=""):
    """向Word文档添加单零件分析章节"""
    topology = analysis["topology"]
    face_types = analysis["face_types"]
    edge_types = analysis["edge_types"]
    bbox = analysis["bbox"]
    volume = analysis["volume"]
    center_of_mass = analysis["center_of_mass"]
    surface_area = analysis["surface_area"]
    holes = analysis["holes"]
    tubes = analysis["tubes"]
    sq_pairs = analysis["sq_pairs"]
    sq_cross = analysis["sq_cross"]
    wall_info = analysis["wall_info"]
    model_type = analysis["model_type"]

    model_type_desc = {
        "管状结构": "该模型包含同轴的内外圆柱面，属于管状结构",
        "方通结构": "该模型包含多组平行平面对，属于方通结构",
        "一般机械件": "该模型为一般机械零件",
    }

    # 1. 概览
    doc.add_heading(f"概览{part_name}", level=1)
    overview_table = doc.add_table(rows=0, cols=2)
    overview_table.style = 'Table Grid'
    overview_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_row(overview_table, ["模型类型", model_type])
    add_table_row(overview_table, ["简要描述", model_type_desc.get(model_type, "")])
    add_table_row(overview_table, ["顶点数", str(topology["顶点"])])
    add_table_row(overview_table, ["边数", str(topology["边"])])
    add_table_row(overview_table, ["面数", str(topology["面"])])

    # 2. 拓扑信息
    doc.add_heading(f"拓扑信息{part_name}", level=1)
    topo_table = doc.add_table(rows=1, cols=2)
    topo_table.style = 'Table Grid'
    topo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    topo_table.rows[0].cells[0].text = "拓扑类型"
    topo_table.rows[0].cells[1].text = "数量"
    for name in ["顶点", "边", "线", "面", "壳", "实体"]:
        add_table_row(topo_table, [name, str(topology[name])])

    # 3. 几何类型分布
    doc.add_heading(f"几何类型分布{part_name}", level=1)

    doc.add_heading("面类型统计", level=2)
    face_table = doc.add_table(rows=1, cols=2)
    face_table.style = 'Table Grid'
    face_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    face_table.rows[0].cells[0].text = "面类型"
    face_table.rows[0].cells[1].text = "数量"
    for tname, count in sorted(face_types.items(), key=lambda x: -x[1]):
        add_table_row(face_table, [tname, str(count)])

    doc.add_heading("边类型统计", level=2)
    edge_table = doc.add_table(rows=1, cols=2)
    edge_table.style = 'Table Grid'
    edge_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    edge_table.rows[0].cells[0].text = "边类型"
    edge_table.rows[0].cells[1].text = "数量"
    for tname, count in sorted(edge_types.items(), key=lambda x: -x[1]):
        add_table_row(edge_table, [tname, str(count)])

    # 4. 尺寸信息
    doc.add_heading(f"尺寸信息{part_name}", level=1)
    dim_table = doc.add_table(rows=0, cols=2)
    dim_table.style = 'Table Grid'
    dim_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_row(dim_table, ["X方向尺寸", fmt_length(bbox["x_size"])])
    add_table_row(dim_table, ["Y方向尺寸", fmt_length(bbox["y_size"])])
    add_table_row(dim_table, ["Z方向尺寸", fmt_length(bbox["z_size"])])
    add_table_row(dim_table, ["X范围", f"[{bbox['x_min']:.3f}, {bbox['x_max']:.3f}] mm"])
    add_table_row(dim_table, ["Y范围", f"[{bbox['y_min']:.3f}, {bbox['y_max']:.3f}] mm"])
    add_table_row(dim_table, ["Z范围", f"[{bbox['z_min']:.3f}, {bbox['z_max']:.3f}] mm"])
    add_table_row(dim_table, ["体积", fmt_volume(volume)])
    add_table_row(dim_table, ["表面积", fmt_area(surface_area)])
    com_x, com_y, com_z = center_of_mass.X(), center_of_mass.Y(), center_of_mass.Z()
    add_table_row(dim_table, ["质心位置", f"({com_x:.3f}, {com_y:.3f}, {com_z:.3f}) mm"])

    # 5. 孔特征
    doc.add_heading(f"孔特征{part_name}", level=1)
    if holes:
        regular_count = sum(1 for h in holes if h["classification"].startswith("规则"))
        irregular_count = sum(1 for h in holes if h["classification"].startswith("不规则"))
        doc.add_paragraph(f"孔总数: {len(holes)}，规则孔: {regular_count}，不规则孔: {irregular_count}")

        hole_table = doc.add_table(rows=1, cols=7)
        hole_table.style = 'Table Grid'
        hole_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["序号", "分类", "直径(mm)", "半径(mm)", "槽长(mm)", "槽宽(mm)", "轴位置/备注"]
        for i, h in enumerate(headers):
            hole_table.rows[0].cells[i].text = h

        for idx, hole in enumerate(holes, 1):
            diameter = f"{hole.get('diameter', 0):.3f}" if hole.get('diameter') else "-"
            radius = f"{hole.get('radius', 0):.3f}" if hole.get('radius') else "-"
            slot_length = f"{hole.get('slot_length', 0):.3f}" if hole.get('slot_length') else "-"
            slot_width = f"{hole.get('slot_width', 0):.3f}" if hole.get('slot_width') else "-"
            axis_loc = hole.get("axis_location", None)
            axis_dir = hole.get("axis_direction", None)
            if axis_loc:
                note_str = f"({axis_loc[0]:.3f}, {axis_loc[1]:.3f}, {axis_loc[2]:.3f})"
            elif axis_dir:
                note_str = f"方向({axis_dir[0]:.3f}, {axis_dir[1]:.3f}, {axis_dir[2]:.3f})"
            else:
                note_str = "-"
            add_table_row(hole_table, [
                str(idx), hole["classification"], diameter, radius, slot_length, slot_width, note_str
            ])
    else:
        doc.add_paragraph("未检测到孔特征")

    # 6. 管状/方通结构
    if tubes:
        doc.add_heading(f"管状结构{part_name}", level=1)
        doc.add_paragraph(f"检测到 {len(tubes)} 组同轴圆柱面对")

        tube_table = doc.add_table(rows=1, cols=6)
        tube_table.style = 'Table Grid'
        tube_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["序号", "外径(mm)", "内径(mm)", "壁厚(mm)", "长度(mm)", "轴方向"]
        for i, h in enumerate(headers):
            tube_table.rows[0].cells[i].text = h

        for idx, tube in enumerate(tubes, 1):
            d = tube["axis_direction"]
            dir_str = f"({d[0]:.3f}, {d[1]:.3f}, {d[2]:.3f})"
            add_table_row(tube_table, [
                str(idx),
                fmt_length(tube["outer_diameter"]),
                fmt_length(tube["inner_diameter"]),
                fmt_length(tube["wall_thickness"]),
                fmt_length(tube["length"]),
                dir_str,
            ])

    if sq_cross:
        doc.add_heading(f"方通结构{part_name}", level=1)
        for idx, sq in enumerate(sq_cross, 1):
            doc.add_paragraph(
                f"方通截面 {idx}: "
                f"尺寸1 = {fmt_length(sq['cross_section_dim1'])}, "
                f"尺寸2 = {fmt_length(sq['cross_section_dim2'])}, "
                f"方向1对数 = {sq['pairs_count_1']}, "
                f"方向2对数 = {sq['pairs_count_2']}"
            )
    elif sq_pairs:
        doc.add_heading(f"方通结构{part_name}", level=1)
        doc.add_paragraph(f"检测到 {len(sq_pairs)} 组平行平面对")
        sq_table = doc.add_table(rows=1, cols=4)
        sq_table.style = 'Table Grid'
        sq_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["序号", "法向量方向", "间距(mm)", "法向量是否相反"]
        for i, h in enumerate(headers):
            sq_table.rows[0].cells[i].text = h
        for idx, pair in enumerate(sq_pairs, 1):
            d = pair["direction"]
            dir_str = f"({d[0]:.3f}, {d[1]:.3f}, {d[2]:.3f})"
            add_table_row(sq_table, [
                str(idx), dir_str, fmt_length(pair["distance"]),
                "是" if pair["is_opposite_normal"] else "否"
            ])

    # 7. 壁厚信息
    doc.add_heading(f"壁厚信息{part_name}", level=1)
    if wall_info:
        wall_table = doc.add_table(rows=1, cols=4)
        wall_table.style = 'Table Grid'
        wall_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["序号", "类型", "壁厚(mm)", "备注"]
        for i, h in enumerate(headers):
            wall_table.rows[0].cells[i].text = h
        for idx, wi in enumerate(wall_info, 1):
            note = ""
            if wi.get("outer_radius"):
                note = f"外径={wi['outer_radius']*2:.3f}, 内径={wi['inner_radius']*2:.3f}"
            add_table_row(wall_table, [
                str(idx), wi["type"], fmt_length(wi["wall_thickness"]), note
            ])
    else:
        doc.add_paragraph("未检测到壁厚信息")


def _generate_assembly_report(doc, shape, filename, basename):
    """生成装配件（多Solid Compound）分析报告"""
    solids = decompose_compound(shape)
    print(f"  检测到装配件，包含 {len(solids)} 个零件")

    # 整体分析
    overall_topology = count_topology(shape)
    overall_bbox = get_bounding_box(shape)
    overall_volume, overall_com = get_volume(shape)
    overall_surface_area = get_surface_area(shape)

    # 装配件连接性分析
    connectivity = analyze_assembly_connectivity(solids)

    # 每个零件的详细分析
    part_analyses = []
    for idx, solid in enumerate(solids):
        print(f"  分析零件 {idx + 1}/{len(solids)}...")
        analysis = _analyze_single_solid(solid)
        part_analyses.append(analysis)

    # ---- 装配件概览 ----
    doc.add_heading("装配件概览", level=1)
    overview_table = doc.add_table(rows=0, cols=2)
    overview_table.style = 'Table Grid'
    overview_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_row(overview_table, ["文件名", filename])
    add_table_row(overview_table, ["结构类型", f"装配件（{len(solids)} 个零件）"])
    add_table_row(overview_table, ["整体顶点数", str(overall_topology["顶点"])])
    add_table_row(overview_table, ["整体边数", str(overall_topology["边"])])
    add_table_row(overview_table, ["整体面数", str(overall_topology["面"])])
    add_table_row(overview_table, ["整体X尺寸", fmt_length(overall_bbox["x_size"])])
    add_table_row(overview_table, ["整体Y尺寸", fmt_length(overall_bbox["y_size"])])
    add_table_row(overview_table, ["整体Z尺寸", fmt_length(overall_bbox["z_size"])])
    add_table_row(overview_table, ["整体体积", fmt_volume(overall_volume)])
    add_table_row(overview_table, ["整体表面积", fmt_area(overall_surface_area)])

    # ---- 零件清单 ----
    doc.add_heading("零件清单", level=1)
    part_list_table = doc.add_table(rows=1, cols=5)
    part_list_table.style = 'Table Grid'
    part_list_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["零件编号", "模型类型", "体积(mm³)", "壁厚(mm)", "孔数量"]
    for i, h in enumerate(headers):
        part_list_table.rows[0].cells[i].text = h

    for idx, (analysis, conn) in enumerate(zip(part_analyses, connectivity)):
        wall_thick = "-"
        if analysis["wall_info"]:
            wall_thick = fmt_length(analysis["wall_info"][0]["wall_thickness"])
        elif analysis["tubes"]:
            wall_thick = fmt_length(analysis["tubes"][0]["wall_thickness"])
        add_table_row(part_list_table, [
            f"零件 {idx + 1}",
            analysis["model_type"],
            fmt_volume(analysis["volume"]),
            wall_thick,
            str(len(analysis["holes"])),
        ])

    # ---- 端面连接性分析 ----
    doc.add_heading("端面连接性分析", level=1)
    doc.add_paragraph("判断每个零件的端面是自由端（切口）还是焊接端（与其他零件连接）")

    conn_table = doc.add_table(rows=1, cols=5)
    conn_table.style = 'Table Grid'
    conn_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["零件编号", "端面位置", "状态", "连接对象", "说明"]
    for i, h in enumerate(headers):
        conn_table.rows[0].cells[i].text = h

    total_free_ends = 0
    for idx, conn in enumerate(connectivity):
        if not conn["end_faces"]:
            add_table_row(conn_table, [
                f"零件 {idx + 1}", "-", "无端面", "-", "非管/方通结构或端面未识别"
            ])
            continue
        for ef in conn["end_faces"]:
            pos_name = "起始端" if ef["position"] == "min" else "末端"
            if ef["is_free"]:
                status = "自由端（切口）"
                total_free_ends += 1
                connected = "-"
                desc = "该端面未被其他零件覆盖，属于切口"
            else:
                status = "焊接端"
                connected = f"零件 {ef['connected_to'] + 1}"
                desc = "该端面与其他零件连接，不属于切口"
            add_table_row(conn_table, [
                f"零件 {idx + 1}", pos_name, status, connected, desc
            ])

    doc.add_paragraph(f"自由端（切口）总数: {total_free_ends}")

    # ---- 每个零件详细分析 ----
    for idx, analysis in enumerate(part_analyses):
        doc.add_heading(f"零件 {idx + 1} 详细分析", level=1)
        _add_single_part_sections(doc, analysis, part_name=f"（零件 {idx + 1}）")

    # ---- 装配件汇总 ----
    doc.add_heading("装配件汇总", level=1)
    total_holes = sum(len(a["holes"]) for a in part_analyses)
    regular_holes = sum(sum(1 for h in a["holes"] if h["classification"].startswith("规则")) for a in part_analyses)
    irregular_holes = total_holes - regular_holes
    total_tubes = sum(len(a["tubes"]) for a in part_analyses)

    summary_table = doc.add_table(rows=0, cols=2)
    summary_table.style = 'Table Grid'
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_row(summary_table, ["零件总数", str(len(solids))])
    add_table_row(summary_table, ["自由端（切口）总数", str(total_free_ends)])
    add_table_row(summary_table, ["孔总数", str(total_holes)])
    add_table_row(summary_table, ["规则孔数", str(regular_holes)])
    add_table_row(summary_table, ["不规则孔数", str(irregular_holes)])
    add_table_row(summary_table, ["管状结构组数", str(total_tubes)])
    add_table_row(summary_table, ["整体体积", fmt_volume(overall_volume)])


def generate_report(filepath, output_dir, is_assembly=False):
    """生成Word分析报告

    Args:
        filepath: STEP文件路径
        output_dir: 输出目录
        is_assembly: 是否为装配件（True=多零件装配件，False=单零件）
    """
    filename = os.path.basename(filepath)
    basename = os.path.splitext(filename)[0]

    print(f"正在分析: {filename}")
    print(f"  分析模式: {'装配件' if is_assembly else '单零件'}")

    # 读取STEP文件
    shape = read_step_file(filepath)
    if shape is None:
        print(f"  错误: 无法读取文件 {filename}")
        return False

    # 创建Word文档
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(10.5)

    # 标题
    if is_assembly:
        doc.add_heading(f"{basename} 装配件分析报告", level=0)
    else:
        doc.add_heading(f"{basename} 分析报告", level=0)

    # 根据模式走不同路径
    if is_assembly:
        _generate_assembly_report(doc, shape, filename, basename)
    else:
        # 单零件分析
        analysis = _analyze_single_solid(shape)
        # 概览中添加文件名
        overview_table = doc.add_table(rows=0, cols=2)
        overview_table.style = 'Table Grid'
        overview_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_table_row(overview_table, ["文件名", filename])
        model_type_desc = {
            "管状结构": "该模型包含同轴的内外圆柱面，属于管状结构",
            "方通结构": "该模型包含多组平行平面对，属于方通结构",
            "一般机械件": "该模型为一般机械零件",
        }
        add_table_row(overview_table, ["模型类型", analysis["model_type"]])
        add_table_row(overview_table, ["简要描述", model_type_desc.get(analysis["model_type"], "")])
        add_table_row(overview_table, ["顶点数", str(analysis["topology"]["顶点"])])
        add_table_row(overview_table, ["边数", str(analysis["topology"]["边"])])
        add_table_row(overview_table, ["面数", str(analysis["topology"]["面"])])

        # 其余章节
        _add_single_part_sections(doc, analysis, part_name="")

    # 保存文档
    output_filename = f"{basename}_分析报告.docx"
    output_path = os.path.join(output_dir, output_filename)
    doc.save(output_path)
    print(f"  报告已生成: {output_path}")
    return True


# ============================================================
# 主函数
# ============================================================
def main():
    parser = argparse.ArgumentParser(description="STEP文件分析工具 - 解析STEP文件并生成Word分析报告")
    parser.add_argument("--input", required=True, help="STEP文件路径或包含STEP文件的目录")
    parser.add_argument("--output", help="输出目录，默认与输入相同")
    parser.add_argument("--assembly", action="store_true", default=False,
                        help="指定为装配件模式（多零件Compound），默认为单零件模式")
    args = parser.parse_args()

    input_path = args.input
    output_dir = args.output
    is_assembly = args.assembly

    # 收集STEP文件
    step_files = []
    if os.path.isfile(input_path):
        step_files.append(input_path)
        if output_dir is None:
            output_dir = os.path.dirname(os.path.abspath(input_path))
    elif os.path.isdir(input_path):
        if output_dir is None:
            output_dir = input_path
        for f in os.listdir(input_path):
            ext = os.path.splitext(f)[1].lower()
            if ext in ('.step', '.stp'):
                step_files.append(os.path.join(input_path, f))
    else:
        print(f"错误: 输入路径不存在: {input_path}")
        sys.exit(1)

    if not step_files:
        print("未找到STEP文件")
        sys.exit(1)

    # 确保输出目录存在
    os.makedirs(output_dir, exist_ok=True)

    # 处理每个文件
    success_count = 0
    fail_count = 0
    for step_file in step_files:
        try:
            result = generate_report(step_file, output_dir, is_assembly=is_assembly)
            if result:
                success_count += 1
            else:
                fail_count += 1
        except Exception as e:
            print(f"  错误: 处理文件 {os.path.basename(step_file)} 时出错: {e}")
            fail_count += 1

    print(f"\n处理完成: 成功 {success_count} 个, 失败 {fail_count} 个")


if __name__ == "__main__":
    main()
